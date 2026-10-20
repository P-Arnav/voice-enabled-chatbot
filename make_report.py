"""
Build the brief Word report for the Voice-Enabled Chatbot lab assessment.

Reads  : results*.json, figures/*.png
Writes : ../Voice_Chatbot_Lab_Assessment.docx

Run:  python make_report.py
"""

import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAME = "Arnav Panwar"
REG_NO = "23BAI0147"
LIVE_URL = "LIVE_URL_PENDING"
REPO_URL = "https://github.com/P-Arnav/voice-enabled-chatbot"
MODEL_URL = "https://huggingface.co/Xyphitos/distilbert-clinc25-intent"
OUT_FILE = "../Voice_Chatbot_Lab_Assessment.docx"

bil = json.load(open("results.json"))
dis = json.load(open("results_distilbert.json"))
voice = json.load(open("results_voice.json"))

doc = Document()
for s in doc.sections:
    s.page_height, s.page_width = Inches(11.69), Inches(8.27)          # A4
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.8)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for h in ("Heading 1", "Heading 2"):
    doc.styles[h].font.name = "Calibri"
    doc.styles[h].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def para(text, bold=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)


def hyperlink(paragraph, url, text=None):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/"
                                   "2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "1F5FBF")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text or url
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def table(rows, widths, bold_last=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(val)
            cell.width = Inches(widths[j])
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.bold = i == 0 or (bold_last and i == len(rows) - 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(path, caption, width=6.3):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(caption, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    p.runs[0].italic = True


def pct(x):
    return f"{x * 100:.1f} %"


# footer
fp = doc.sections[0].footer.paragraphs[0]
fp.text = f"{NAME}  |  {REG_NO}"
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------------------------------------------------------------------------
para("Lab Assessment", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para("Voice-Enabled Chatbot using Speech Recognition and Deep Learning", bold=True,
     size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(f"{NAME}  ({REG_NO})", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

for label, url in [("Live application: ", LIVE_URL), ("Source code: ", REPO_URL),
                   ("Trained model: ", MODEL_URL)]:
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    hyperlink(p, url)
    p.paragraph_format.space_after = Pt(2)

doc.add_heading("1. Objective", level=1)
para("The aim is to build and publicly deploy a chatbot that users can talk to. The "
     "application records the user's voice in the browser, converts the speech to text "
     "with a deep-learning speech recognition model, classifies the intent of the text "
     "with a deep-learning intent classifier, and generates a reply. The page shows the "
     "recognized speech, the predicted intent with its confidence, and the chatbot "
     "response, and it also speaks the response aloud.")

doc.add_heading("2. System Overview", level=1)
figure("figures/pipeline.png", "Figure 1: End-to-end pipeline of the voice-enabled chatbot.")
bullet(" the browser microphone widget (st.audio_input) records a WAV clip.", "Voice input:")
bullet(" OpenAI Whisper (base.en) transcribes the clip to text.", "Speech recognition:")
bullet(" a fine-tuned DistilBERT model maps the text to one of 25 intents or "
       "'out-of-scope'.", "Intent classification:")
bullet(" a template reply is selected, or a dynamic one is computed for time, date, "
       "arithmetic, spelling, dictionary definitions, coin flips and dice rolls.",
       "Response generation:")
bullet(" the reply is shown as text, spoken with gTTS and added to the chat history.",
       "Output:")

doc.add_heading("3. Dataset", level=1)
para("The intent model is trained on CLINC150 (Larson et al., 2019), a public intent "
     "classification benchmark. It contains 150 intents across 10 domains, with 100 "
     "training, 20 validation and 30 test utterances per intent, plus a dedicated set of "
     "out-of-scope (OOS) queries. Out-of-scope queries matter for a real chatbot, because "
     "users often ask things the bot cannot handle.")
para("Many CLINC150 intents (banking, travel, car maintenance and so on) need back-end "
     "services that this bot does not have. I therefore selected 25 intents the bot can "
     "actually answer, covering small talk, the bot's identity, time and date, and "
     "utilities. The other 125 intents are used as additional negative examples for the "
     "OOS class, so the model learns to reject requests it cannot serve.")
table([
    ["Split", "In-scope (25 intents)", "Out-of-scope", "Total"],
    ["Train", "2,500", f"{bil['n_train'] - 2500:,} (100 CLINC OOS + 1,500 other intents)",
     f"{bil['n_train']:,}"],
    ["Validation", "500", f"{bil['n_val'] - 500} (100 CLINC OOS + 300 other intents)",
     f"{bil['n_val']}"],
    ["Test", f"{bil['n_test_in']}", f"{bil['n_test_oos']:,} (CLINC OOS test set)",
     f"{bil['n_test_in'] + bil['n_test_oos']:,}"],
], [1.1, 1.6, 2.9, 0.9])
para("Intents: greeting, goodbye, thank_you, yes, no, repeat, time, date, calculator, "
     "flip_coin, roll_dice, spelling, definition, tell_joke, fun_fact, meaning_of_life, "
     "what_is_your_name, who_made_you, how_old_are_you, are_you_a_bot, "
     "what_can_i_ask_you, what_are_your_hobbies, where_are_you_from, do_you_have_pets, "
     "who_do_you_work_for (+ oos).", size=10)

doc.add_heading("4. Model Architecture", level=1)
doc.add_heading("4.1 Speech recognition: Whisper", level=2)
para("Whisper is a Transformer encoder-decoder trained on 680,000 hours of weakly "
     "supervised speech. The audio is resampled to 16 kHz and converted to an 80-channel "
     "log-Mel spectrogram. The encoder processes the spectrogram, and the decoder "
     "generates the transcript token by token. I use the English-only base model "
     "(74M parameters) because it is accurate and fast enough on a free CPU server.")
doc.add_heading("4.2 Intent classification: DistilBERT", level=2)
para("DistilBERT is a 6-layer Transformer encoder (hidden size 768, 12 attention heads, "
     "66M parameters) distilled from BERT-base. The sentence is split into WordPiece "
     "tokens with a [CLS] token added at the start. The final hidden state of [CLS] goes "
     "through a pre-classifier dense layer with ReLU and dropout, then a linear layer with "
     "26 outputs and a softmax. If the highest probability is below 0.5, the prediction "
     "is changed to 'oos' and the bot replies with a fallback message.")
doc.add_heading("4.3 Comparison models", level=2)
bullet(f" Embedding(100) -> bidirectional LSTM(128) -> masked max-pooling over time -> "
       f"Dropout(0.3) -> Linear(26), trained from scratch ({bil['n_params']:,} "
       f"parameters).", "BiLSTM:")
bullet(" word unigram and bigram TF-IDF features with a logistic regression "
       "classifier (classical, non-deep-learning baseline).", "TF-IDF + Logistic Regression:")

doc.add_heading("5. Methodology", level=1)
bullet(" lowercasing and tokenization (WordPiece for DistilBERT, word-level vocabulary "
       f"of {bil['vocab_size']:,} words for the BiLSTM, padded to 25 tokens).",
       "Pre-processing:")
bullet(f" AdamW, learning rate 5e-5 with linear decay, weight decay 0.01, batch size 32, "
       f"{dis['epochs']} epochs, max length 32 tokens. Training took {dis['train_time_s']:.0f} s "
       f"on an RTX GPU.", "DistilBERT fine-tuning:")
bullet(f" Adam, learning rate 1e-3, batch 64, {bil['epochs']} epochs, cross-entropy "
       f"weighted by inverse class frequency (OOS is about 40 % of the training data). "
       f"The checkpoint with the best validation accuracy (epoch {bil['best_epoch']}) was kept.",
       "BiLSTM training:")
bullet(" overall accuracy, in-scope accuracy, OOS recall and macro-F1 on the CLINC150 "
       "test set. For the speech side, 70 test sentences (2 per intent + 20 OOS) were "
       "synthesized with two Windows voices (male and female) and passed through the "
       "full pipeline to measure word error rate (WER) and end-to-end intent accuracy.",
       "Evaluation:")
bullet(" the app is written in Streamlit and hosted on Streamlit Community Cloud from "
       "the GitHub repository. The fine-tuned model is published on the Hugging Face "
       "Hub, and the app downloads it at start-up.", "Deployment:")

doc.add_heading("6. Results", level=1)
doc.add_heading("6.1 Intent classification (CLINC150 test set)", level=2)
b, bt, base, d = bil["bilstm"], bil["bilstm_threshold"], bil["tfidf_logreg"], dis["distilbert"]
table([
    ["Model", "Overall acc", "In-scope acc", "OOS recall", "Macro-F1"],
    ["TF-IDF + LogReg (baseline)", pct(base["overall_acc"]), pct(base["in_scope_acc"]),
     pct(base["oos_recall"]), f"{base['macro_f1']:.3f}"],
    ["BiLSTM (argmax)", pct(b["overall_acc"]), pct(b["in_scope_acc"]),
     pct(b["oos_recall"]), f"{b['macro_f1']:.3f}"],
    ["BiLSTM (+ threshold 0.5)", pct(bt["overall_acc"]), pct(bt["in_scope_acc"]),
     pct(bt["oos_recall"]), f"{bt['macro_f1']:.3f}"],
    ["DistilBERT (+ threshold 0.5), deployed", pct(d["overall_acc"]), pct(d["in_scope_acc"]),
     pct(d["oos_recall"]), f"{d['macro_f1']:.3f}"],
], [2.6, 1.0, 1.0, 1.0, 0.9], bold_last=True)
para(f"Fine-tuned DistilBERT is the best model on every metric. It reaches "
     f"{pct(d['in_scope_acc'])} in-scope accuracy and {pct(d['oos_recall'])} OOS recall, "
     f"with a CPU latency of about {dis['cpu_latency_ms']:.0f} ms per query. The BiLSTM "
     f"trained from scratch rejects OOS queries well but reaches only "
     f"{pct(bt['in_scope_acc'])} in-scope accuracy, which is below the TF-IDF baseline. "
     f"With only 100 examples per intent, the knowledge DistilBERT gained from "
     f"pre-training makes the biggest difference, so DistilBERT is the deployed model.")
figure("figures/distilbert_curves.png", "Figure 2: DistilBERT fine-tuning loss and "
       "validation accuracy.", width=6.0)
figure("figures/confusion_matrix_distilbert.png", "Figure 3: Row-normalised confusion "
       "matrix of the deployed DistilBERT model on the test set.", width=4.4)

doc.add_heading("6.2 End-to-end voice evaluation", level=2)
table([
    ["Metric", "Value"],
    ["Test clips (2 voices)", f"{voice['n_clips']}"],
    ["Whisper word error rate", pct(voice["wer"])],
    ["Intent accuracy from typed text", pct(voice["intent_acc_text"])],
    ["Intent accuracy from voice (end-to-end)", pct(voice["intent_acc_voice"])],
    ["Speech recognition time per clip (GPU)", f"{voice['asr_seconds_per_clip']:.2f} s"],
], [3.4, 1.4])
para("Whisper's few errors were mostly punctuation, casing and rare words (for example, a "
     "misspelled source word 'politicsu' was heard as 'Politiksoo'). None of them changed "
     "the predicted intent, so voice input gives the same intent accuracy as typed input "
     "on this test. The only error in the set was a classification mistake: 'what are you "
     "programmed to do for me' was labelled who_made_you instead of what_can_i_ask_you.")

doc.add_heading("6.3 Deployed application", level=2)
if os.path.exists("figures/app_screenshot.png"):
    figure("figures/app_screenshot.png", "Figure 4: The deployed application showing the "
           "recognized speech, detected intent and chatbot response.", width=6.3)
else:
    para("[Screenshot of the live application to be inserted]")

doc.add_heading("7. Conclusion", level=1)
para("The voice-enabled chatbot combines Whisper speech recognition with a fine-tuned "
     "DistilBERT intent classifier and is publicly deployed on Streamlit Community Cloud. "
     "The deep-learning intent model clearly outperforms a classical TF-IDF baseline and "
     "a BiLSTM trained from scratch, and it rejects out-of-scope requests reliably. "
     "Possible extensions are support for more intents and back-end services (weather, "
     "reminders), multilingual speech with the multilingual Whisper models, and testing "
     "with recordings of real users rather than synthesized speech.")

doc.add_heading("References", level=1)
for ref in [
    "S. Larson et al., \"An Evaluation Dataset for Intent Classification and Out-of-Scope "
    "Prediction,\" EMNLP-IJCNLP, 2019.",
    "A. Radford et al., \"Robust Speech Recognition via Large-Scale Weak Supervision,\" "
    "ICML, 2023.",
    "V. Sanh et al., \"DistilBERT, a distilled version of BERT: smaller, faster, cheaper "
    "and lighter,\" arXiv:1910.01108, 2019.",
]:
    p = doc.add_paragraph(ref, style="List Number")
    p.runs[0].font.size = Pt(10)

doc.save(OUT_FILE)
print("Saved", OUT_FILE)
